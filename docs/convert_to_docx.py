#!/usr/bin/env python3
"""
Convert Markdown documents to beautifully styled Word (.docx) documents.
Designed for Secure Agentic BI Interim Report & Slide Presentation.
"""

import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color palette: Academic Navy & Teal
COLOR_PRIMARY = RGBColor(15, 28, 46)     # Deep Navy
COLOR_SECONDARY = RGBColor(13, 148, 136) # Teal
COLOR_DARK = RGBColor(30, 41, 59)        # Slate Dark
COLOR_MUTED = RGBColor(100, 116, 139)    # Slate Muted
COLOR_HIGHLIGHT = RGBColor(194, 59, 59)  # Deep Red for Alerts
HEX_TABLE_HEADER = "0F1C2E"              # Table header background
HEX_TABLE_ALT = "F8FAFC"                 # Table alternate row
HEX_BORDER = "CBD5E1"                    # Table border
HEX_CODE_BG = "F1F5F9"                   # Code block background
HEX_QUOTE_BG = "F0FDF4"                  # Quote background

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table):
    """Apply elegant, subtle borders to table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_TABLE_HEADER}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def format_inline_runs(paragraph, text, base_font="Calibri", base_size=Pt(11), is_bold_default=False, is_italic_default=False, default_color=COLOR_DARK):
    """
    Parse inline bold (**text**), italic (*text*), and code (`text`) formatting.
    """
    # Regex to capture bold-italic, bold, italic, code
    token_pattern = re.compile(
        r'(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    )
    parts = token_pattern.split(text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = base_font
        run.font.size = base_size
        run.font.color.rgb = default_color
        run.bold = is_bold_default
        run.italic = is_italic_default

        if part.startswith('***') and part.endswith('***') and len(part) >= 6:
            run.text = part[3:-3]
            run.bold = True
            run.italic = True
        elif part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt * 0.92)
            run.font.color.rgb = COLOR_SECONDARY
        else:
            run.text = part

def build_docx_from_markdown(md_file_path, output_docx_path, doc_type="report"):
    """Parse markdown file and write styled docx file."""
    doc = Document()

    # Configure Margins: 1 inch on all sides
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = COLOR_DARK

    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for l in tbl_lines:
            cleaned = l.strip().strip('|')
            cells = [c.strip() for c in cleaned.split('|')]
            if not any(set(c).issubset({'-', ':', ' '}) for c in cells):
                rows_data.append(cells)
        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        for row_idx, r_data in enumerate(rows_data):
            row = table.rows[row_idx]
            is_header = (row_idx == 0)
            for col_idx in range(cols_count):
                cell = row.cells[col_idx]
                cell_text = r_data[col_idx] if col_idx < len(r_data) else ""
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.1

                if is_header:
                    set_cell_background(cell, HEX_TABLE_HEADER)
                    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                    format_inline_runs(p, cell_text, base_font="Calibri", base_size=Pt(10), is_bold_default=True, default_color=RGBColor(255, 255, 255))
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, HEX_TABLE_ALT)
                    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
                    format_inline_runs(p, cell_text, base_font="Calibri", base_size=Pt(9.5), default_color=COLOR_DARK)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(4)
        p_spacer.paragraph_format.space_after = Pt(6)

    def flush_code_block(c_lines):
        if not c_lines:
            return
        # Render code block inside a single-cell shaded table for clean bordered look
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        # Subtle border for code block
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_TABLE_HEADER}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'</w:tcBorders>'
        )
        tc_pr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        code_text = "".join(c_lines)
        run = p.add_run(code_text.rstrip())
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_DARK

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(2)
        p_spacer.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')

        # Check code fence
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block(code_block_lines)
                code_block_lines = []
            else:
                if table_lines:
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(raw_line)
            i += 1
            continue

        # Check table line
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line)
            i += 1
            continue
        else:
            if table_lines:
                flush_table(table_lines)
                table_lines = []

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Horizontal Rule
        if re.match(r'^-{3,}$', line.strip()):
            p_hr = doc.add_paragraph()
            p_hr.paragraph_format.space_before = Pt(6)
            p_hr.paragraph_format.space_after = Pt(6)
            p_hr_border = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="{HEX_BORDER}"/>'
                f'</w:pBdr>'
            )
            p_hr._p.get_or_add_pPr().append(p_hr_border)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            text = line[2:].strip()
            format_inline_runs(p, text, base_font="Calibri", base_size=Pt(20), is_bold_default=True, default_color=COLOR_PRIMARY)
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            text = line[3:].strip()
            format_inline_runs(p, text, base_font="Calibri", base_size=Pt(15), is_bold_default=True, default_color=COLOR_PRIMARY)
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            text = line[4:].strip()
            format_inline_runs(p, text, base_font="Calibri", base_size=Pt(12.5), is_bold_default=True, default_color=COLOR_SECONDARY)
            i += 1
            continue

        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            text = line[5:].strip()
            format_inline_runs(p, text, base_font="Calibri", base_size=Pt(11.5), is_bold_default=True, default_color=COLOR_DARK)
            i += 1
            continue

        # Blockquote (Speaker notes or Callouts)
        if line.startswith('> '):
            quote_text = line[2:].strip()
            # If next lines also start with >
            while i + 1 < len(lines) and lines[i+1].startswith('> '):
                i += 1
                quote_text += " " + lines[i][2:].strip()

            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "F0FDF4" if "Speaker Notes" in quote_text or "Lý do" in quote_text else "F8FAFC")
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

            tc_pr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{COLOR_SECONDARY.rgb if isinstance(COLOR_SECONDARY.rgb, str) else "0D9488"}"/>'
                f'<w:top w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tc_pr.append(borders)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, quote_text, base_font="Calibri", base_size=Pt(10), is_italic_default=True, default_color=COLOR_DARK)

            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_before = Pt(2)
            p_spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Bullet List Items
        bullet_match = re.match(r'^(\s*)[*+-]\s+(.*)$', line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_text = bullet_match.group(2)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 + (indent_spaces // 2) * 0.2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, bullet_text, base_font="Calibri", base_size=Pt(10.5), default_color=COLOR_DARK)
            i += 1
            continue

        # Numbered List Items
        num_match = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if num_match:
            indent_spaces = len(num_match.group(1))
            num_text = num_match.group(3)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Inches(0.25 + (indent_spaces // 2) * 0.2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, num_text, base_font="Calibri", base_size=Pt(10.5), default_color=COLOR_DARK)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.18
        format_inline_runs(p, line, base_font="Calibri", base_size=Pt(10.5), default_color=COLOR_DARK)
        i += 1

    # Flush any remaining table or code block
    if table_lines:
        flush_table(table_lines)
    if code_block_lines:
        flush_code_block(code_block_lines)

    doc.save(output_docx_path)
    print(f"Successfully generated: {output_docx_path}")

if __name__ == '__main__':
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    docs_dir = os.path.join(project_root, "docs")

    interim_md = os.path.join(docs_dir, "Interim_Report.md")
    interim_docx = os.path.join(docs_dir, "Interim_Report.docx")
    interim_root_docx = os.path.join(project_root, "Interim_Report.docx")

    slide_md = os.path.join(docs_dir, "Slide_Presentation.md")
    slide_docx = os.path.join(docs_dir, "Slide_Presentation.docx")
    slide_root_docx = os.path.join(project_root, "Slide_Presentation.docx")

    print(f"Converting {interim_md} to docx...")
    build_docx_from_markdown(interim_md, interim_docx, doc_type="report")
    build_docx_from_markdown(interim_md, interim_root_docx, doc_type="report")

    print(f"Converting {slide_md} to docx...")
    build_docx_from_markdown(slide_md, slide_docx, doc_type="slide")
    build_docx_from_markdown(slide_md, slide_root_docx, doc_type="slide")

    print("All conversions completed successfully!")
